from __future__ import annotations

from pathlib import Path
from typing import Any

from .db import load_report_bundle
from .paths import export_dir, safe_filename
from .utils import balance_text, format_date, is_full_year, money, to_float


def _period_line(report: dict[str, Any]) -> str:
    try:
        year = int(report.get("year") or 0)
    except ValueError:
        year = 0
    start = report.get("period_start") or ""
    end = report.get("period_end") or ""
    show = bool(report.get("show_period_note"))
    if start and end and (show or not is_full_year(year, start, end)):
        return f"Фактический период управления: с {format_date(start)} по {format_date(end)}"
    return ""




def _management_cost(bundle: dict[str, Any]) -> float:
    # Only the stored official total is exported. A zero can be an explicit
    # managerial decision; internal detail is transferred via the UI's
    # "Посчитать из строк ниже" action, never implicitly at export time.
    return to_float(bundle.get("report", {}).get("management_cost"))


def _financial_rows(bundle: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for fin in bundle["financials"]:
        payer = fin.get("payer_type") or ""
        values = [to_float(fin.get(k)) for k in ("debt_start", "accrued", "paid", "debt_end")]
        if payer == "Наниматели" and all(abs(v) < 0.0001 for v in values) and not (fin.get("notes") or "").strip():
            continue
        rows.append(fin)
    if not rows:
        rows.append({"payer_type": "Собственники", "debt_start": 0, "accrued": 0, "paid": 0, "debt_end": 0})
    return rows


def export_xlsx(report_id: int) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    bundle = load_report_bundle(report_id)
    report = bundle["report"]
    company = bundle["company"]
    house = bundle["house"]
    out_dir = export_dir(company.get("name", "УО"), report.get("year", "год"), house.get("address", "дом"))
    out_path = out_dir / f"Отчет_728_{safe_filename(str(house.get('address','дом')), 'дом', 55)}_{report.get('year')}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет 728пр"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    thin = Side(style="thin", color="333333")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E9EEF7")
    title_font = Font(name="Arial", size=13, bold=True)
    bold = Font(name="Arial", size=10, bold=True)
    normal = Font(name="Arial", size=10)

    for col, width in enumerate([6, 28, 12, 14, 16, 16, 16, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    r = 1

    def merge_text(text: str, size: int = 10, is_bold: bool = False, rows: int = 1):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        cell = ws.cell(r, 1, text)
        cell.font = Font(name="Arial", size=size, bold=is_bold)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.row_dimensions[r].height = max(18, 15 * rows)
        r += 1

    def section(text: str):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        cell = ws.cell(r, 1, text)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        r += 1

    def table(headers: list[str], data: list[list[Any]], widths_count: int):
        nonlocal r
        start = r
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r, c, h)
            cell.font = bold
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        r += 1
        for row_data in data:
            for c, value in enumerate(row_data, 1):
                cell = ws.cell(r, c, value)
                cell.font = normal
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        if not data:
            for c in range(1, widths_count + 1):
                cell = ws.cell(r, c, "")
                cell.border = border
            r += 1
        ws.row_dimensions[start].height = 42

    merge_text("ОТЧЕТ", 14, True)
    merge_text("о деятельности по управлению многоквартирным домом", 12, True)
    merge_text(f"по адресу: {house.get('address','')}", 10, False, 2)
    merge_text(f"за {report.get('year','____')} год", 10, True)
    line = _period_line(report)
    if line:
        merge_text(line, 10, False)
    merge_text(str(company.get("name", "")), 10)
    merge_text(f"Адрес приема населения: {company.get('reception_address','')}", 10)
    merge_text(f"ОГРН/ИНН: {company.get('ogrn_inn','')}", 10)
    merge_text(f"Лицо, уполномоченное давать разъяснения: {company.get('report_person','')}", 10)
    merge_text(str(company.get("report_contact", "")), 10)
    merge_text(f"Общая площадь жилых и нежилых помещений: {money(house.get('total_area',0))} м²", 10)
    merge_text(f"Дата размещения отчета: {format_date(report.get('publish_date')) or '________________'}", 10)
    r += 1

    section("1. За отчетный период выполнены следующие работы (оказаны следующие услуги) по содержанию общего имущества собственников помещений в многоквартирном доме:")
    content_data = []
    total_plan = 0.0
    total_actual = 0.0
    for i, item in enumerate(bundle["content"], 1):
        plan_cost = to_float(item.get("plan_cost"))
        actual_cost = to_float(item.get("actual_cost"))
        total_plan += plan_cost
        total_actual += actual_cost
        content_data.append([i, item.get("name", ""), item.get("unit", ""), money(item.get("unit_price")), money(item.get("plan_qty")), money(plan_cost), money(item.get("actual_qty")), money(actual_cost)])
    content_data.append(["ИТОГО", "", "", "", "", money(total_plan), "", money(total_actual)])
    table(["N п/п", "Наименование работы (услуги)", "Ед. изм.", "Цена единицы, руб.", "Кол-во по перечню", "Стоимость по перечню, руб.", "Выполнено, кол-во", "Стоимость выполнено, руб."], content_data, 8)
    r += 1

    section("2. За отчетный период выполнены следующие работы по текущему ремонту общего имущества собственников помещений в многоквартирном доме:")
    merge_text(f"Остаток (перерасход, сальдо) на 1 января отчетного периода: {balance_text(report.get('repair_start_balance'))}", 10)
    merge_text(f"Общий объем денежных средств, подлежащий внесению собственниками за текущий ремонт: {money(report.get('repair_accrual'))} руб.", 10)
    merge_text(f"Стоимость работ по текущему ремонту, выполненных за отчетный период: {money(report.get('repair_work_cost'))} руб.", 10)
    merge_text(f"Остаток (перерасход, сальдо) на 31 декабря отчетного периода: {balance_text(report.get('repair_end_balance'))}", 10)
    repair_data = []
    total_repair = 0.0
    for i, item in enumerate(bundle["repair"], 1):
        total_repair += to_float(item.get("cost"))
        repair_data.append([i, item.get("name", ""), item.get("basis", ""), money(item.get("cost")), item.get("volume", ""), item.get("act_details", "")])
    repair_data.append(["ИТОГО", "", "", money(total_repair), "", ""])
    table(["N п/п", "Наименование работы", "Основание проведения", "Стоимость, руб.", "Объем выполненных работ", "Реквизиты акта или ссылка"], repair_data, 6)
    r += 1

    section("3. Стоимость услуг по управлению многоквартирным домом, оказанных за отчетный период:")
    merge_text(f"{money(_management_cost(bundle))} руб.", 10, True)
    r += 1

    section("4. Сведения о претензионно-исковой работе:")
    claims = bundle["claims"] or {}
    table(["N п/п", "Количество претензий", "Количество исковых заявлений / судебных приказов", "Сумма поступивших средств, руб."], [[1, claims.get("pretensions_count", 0), claims.get("lawsuits_count", 0), money(claims.get("recovered_amount", 0))]], 4)
    r += 1

    section("5. Сведения о начислениях и поступлении средств за отчетный период:")
    fin_rows = _financial_rows(bundle)
    fin_data = []
    total = {"debt_start": 0.0, "accrued": 0.0, "paid": 0.0, "debt_end": 0.0}
    for i, fin in enumerate(fin_rows, 1):
        name = "Платежи собственников помещений в многоквартирном доме" if fin.get("payer_type") == "Собственники" else "Платежи нанимателей помещений в многоквартирном доме"
        for key in total:
            total[key] += to_float(fin.get(key))
        fin_data.append([i, name, money(fin.get("debt_start")), money(fin.get("accrued")), money(fin.get("paid")), money(fin.get("debt_end"))])
    fin_data.append(["ИТОГО", "", money(total["debt_start"]), money(total["accrued"]), money(total["paid"]), money(total["debt_end"])])
    table(["N п/п", "Вид платежа", "Задолженность на начало периода, руб.", "Размер начисленных средств, руб.", "Размер поступивших средств, руб.", "Задолженность на 1 января следующего периода, руб."], fin_data, 6)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
    wb.save(out_path)
    return out_path


def export_docx(report_id: int) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    bundle = load_report_bundle(report_id)
    report = bundle["report"]
    company = bundle["company"]
    house = bundle["house"]
    out_dir = export_dir(company.get("name", "УО"), report.get("year", "год"), house.get("address", "дом"))
    out_path = out_dir / f"Отчет_728_{safe_filename(str(house.get('address','дом')), 'дом', 55)}_{report.get('year')}.docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    def p(text: str, bold: bool = False, center: bool = False):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        return para

    p("ОТЧЕТ", True, True)
    p("о деятельности по управлению многоквартирным домом", True, True)
    p(f"по адресу: {house.get('address','')}", False, True)
    p(f"за {report.get('year','____')} год", True, True)
    line = _period_line(report)
    if line:
        p(line, False, True)
    p(company.get("name", ""), False, True)
    p(f"Адрес приема населения: {company.get('reception_address','')}")
    p(f"ОГРН/ИНН: {company.get('ogrn_inn','')}")
    p(f"Лицо, уполномоченное давать разъяснения: {company.get('report_person','')}")
    p(str(company.get("report_contact", "")))
    p(f"Общая площадь жилых и нежилых помещений: {money(house.get('total_area',0))} м²")
    p(f"Дата размещения отчета: {format_date(report.get('publish_date')) or '________________'}")

    p("1. За отчетный период выполнены следующие работы (оказаны следующие услуги) по содержанию общего имущества собственников помещений в многоквартирном доме:", True)
    tbl = doc.add_table(rows=1, cols=8)
    tbl.style = "Table Grid"
    for i, h in enumerate(["N п/п", "Наименование работы (услуги)", "Ед. изм.", "Цена единицы, руб.", "Кол-во по перечню", "Стоимость по перечню, руб.", "Выполнено, кол-во", "Стоимость выполнено, руб."]):
        tbl.rows[0].cells[i].text = h
    total_plan = total_actual = 0.0
    for idx, item in enumerate(bundle["content"], 1):
        plan_cost = to_float(item.get("plan_cost"))
        actual_cost = to_float(item.get("actual_cost"))
        total_plan += plan_cost
        total_actual += actual_cost
        vals = [idx, item.get("name", ""), item.get("unit", ""), money(item.get("unit_price")), money(item.get("plan_qty")), money(plan_cost), money(item.get("actual_qty")), money(actual_cost)]
        row = tbl.add_row().cells
        for i, v in enumerate(vals):
            row[i].text = str(v)
    row = tbl.add_row().cells
    row[0].text = "ИТОГО"
    row[5].text = money(total_plan)
    row[7].text = money(total_actual)

    p("2. За отчетный период выполнены следующие работы по текущему ремонту общего имущества собственников помещений в многоквартирном доме:", True)
    p(f"Остаток (перерасход, сальдо) на 1 января отчетного периода: {balance_text(report.get('repair_start_balance'))}")
    p(f"Общий объем денежных средств, подлежащий внесению собственниками за текущий ремонт: {money(report.get('repair_accrual'))} руб.")
    p(f"Стоимость работ по текущему ремонту, выполненных за отчетный период: {money(report.get('repair_work_cost'))} руб.")
    p(f"Остаток (перерасход, сальдо) на 31 декабря отчетного периода: {balance_text(report.get('repair_end_balance'))}")
    tbl = doc.add_table(rows=1, cols=6)
    tbl.style = "Table Grid"
    for i, h in enumerate(["N п/п", "Наименование работы", "Основание проведения", "Стоимость, руб.", "Объем выполненных работ", "Реквизиты акта или ссылка"]):
        tbl.rows[0].cells[i].text = h
    total_repair = 0.0
    for idx, item in enumerate(bundle["repair"], 1):
        total_repair += to_float(item.get("cost"))
        vals = [idx, item.get("name", ""), item.get("basis", ""), money(item.get("cost")), item.get("volume", ""), item.get("act_details", "")]
        row = tbl.add_row().cells
        for i, v in enumerate(vals):
            row[i].text = str(v)
    row = tbl.add_row().cells
    row[0].text = "ИТОГО"
    row[3].text = money(total_repair)

    p("3. Стоимость услуг по управлению многоквартирным домом, оказанных за отчетный период:", True)
    p(f"{money(_management_cost(bundle))} руб.")

    p("4. Сведения о претензионно-исковой работе:", True)
    claims = bundle["claims"] or {}
    tbl = doc.add_table(rows=2, cols=4)
    tbl.style = "Table Grid"
    for i, h in enumerate(["N п/п", "Количество претензий", "Количество исковых заявлений / судебных приказов", "Сумма поступивших средств, руб."]):
        tbl.rows[0].cells[i].text = h
    for i, v in enumerate([1, claims.get("pretensions_count", 0), claims.get("lawsuits_count", 0), money(claims.get("recovered_amount", 0))]):
        tbl.rows[1].cells[i].text = str(v)

    p("5. Сведения о начислениях и поступлении средств за отчетный период:", True)
    tbl = doc.add_table(rows=1, cols=6)
    tbl.style = "Table Grid"
    for i, h in enumerate(["N п/п", "Вид платежа", "Задолженность на начало периода, руб.", "Размер начисленных средств, руб.", "Размер поступивших средств, руб.", "Задолженность на 1 января следующего периода, руб."]):
        tbl.rows[0].cells[i].text = h
    total = {"debt_start": 0.0, "accrued": 0.0, "paid": 0.0, "debt_end": 0.0}
    for idx, fin in enumerate(_financial_rows(bundle), 1):
        name = "Платежи собственников помещений в многоквартирном доме" if fin.get("payer_type") == "Собственники" else "Платежи нанимателей помещений в многоквартирном доме"
        for key in total:
            total[key] += to_float(fin.get(key))
        vals = [idx, name, money(fin.get("debt_start")), money(fin.get("accrued")), money(fin.get("paid")), money(fin.get("debt_end"))]
        row = tbl.add_row().cells
        for i, v in enumerate(vals):
            row[i].text = str(v)
    row = tbl.add_row().cells
    row[0].text = "ИТОГО"
    row[2].text = money(total["debt_start"])
    row[3].text = money(total["accrued"])
    row[4].text = money(total["paid"])
    row[5].text = money(total["debt_end"])
    doc.save(out_path)
    return out_path


def export_pdf(report_id: int) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = "Helvetica"
    for font_path in (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/times.ttf")):
        if font_path.exists():
            font_name = "UKFont"
            pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
            break

    bundle = load_report_bundle(report_id)
    report = bundle["report"]
    company = bundle["company"]
    house = bundle["house"]
    out_dir = export_dir(company.get("name", "УО"), report.get("year", "год"), house.get("address", "дом"))
    out_path = out_dir / f"Отчет_728_{safe_filename(str(house.get('address','дом')), 'дом', 55)}_{report.get('year')}.pdf"

    doc = SimpleDocTemplate(str(out_path), pagesize=A4, leftMargin=12*mm, rightMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("NormalRU", parent=styles["Normal"], fontName=font_name, fontSize=8.5, leading=10, alignment=TA_LEFT)
    title = ParagraphStyle("TitleRU", parent=normal, fontName=font_name, fontSize=12, leading=14, alignment=TA_CENTER, spaceAfter=4)
    section = ParagraphStyle("SectionRU", parent=normal, fontName=font_name, fontSize=9, leading=11, spaceBefore=6, spaceAfter=4)

    elems = []
    P = lambda txt, st=normal: Paragraph(str(txt).replace("\n", "<br/>"), st)
    elems += [P("ОТЧЕТ", title), P("о деятельности по управлению многоквартирным домом", title), P(f"по адресу: {house.get('address','')}", title), P(f"за {report.get('year','____')} год", title)]
    line = _period_line(report)
    if line:
        elems.append(P(line, title))
    elems += [P(company.get("name", ""), title), P(f"Адрес приема населения: {company.get('reception_address','')}", normal), P(f"ОГРН/ИНН: {company.get('ogrn_inn','')}", normal), P(f"Лицо, уполномоченное давать разъяснения: {company.get('report_person','')}", normal), P(str(company.get("report_contact", "")), normal), P(f"Общая площадь жилых и нежилых помещений: {money(house.get('total_area',0))} м²", normal), P(f"Дата размещения отчета: {format_date(report.get('publish_date')) or '________________'}", normal), Spacer(1, 4)]

    def add_table(data, col_widths=None):
        table = Table([[P(x) for x in row] for row in data], colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9eef7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
        ]))
        elems.append(table)
        elems.append(Spacer(1, 5))

    elems.append(P("1. За отчетный период выполнены следующие работы (оказаны следующие услуги) по содержанию общего имущества собственников помещений в многоквартирном доме:", section))
    data = [["N", "Наименование", "Ед.", "Цена", "Кол-во", "Стоимость", "Выполнено", "Стоимость"]]
    total_plan = total_actual = 0.0
    for idx, item in enumerate(bundle["content"], 1):
        plan_cost = to_float(item.get("plan_cost"))
        actual_cost = to_float(item.get("actual_cost"))
        total_plan += plan_cost
        total_actual += actual_cost
        data.append([idx, item.get("name", ""), item.get("unit", ""), money(item.get("unit_price")), money(item.get("plan_qty")), money(plan_cost), money(item.get("actual_qty")), money(actual_cost)])
    data.append(["ИТОГО", "", "", "", "", money(total_plan), "", money(total_actual)])
    add_table(data, [10*mm, 54*mm, 13*mm, 20*mm, 20*mm, 24*mm, 20*mm, 25*mm])

    elems.append(P("2. За отчетный период выполнены следующие работы по текущему ремонту общего имущества собственников помещений в многоквартирном доме:", section))
    elems.append(P(f"Остаток (перерасход, сальдо) на 1 января отчетного периода: {balance_text(report.get('repair_start_balance'))}"))
    elems.append(P(f"Общий объем денежных средств, подлежащий внесению собственниками за текущий ремонт: {money(report.get('repair_accrual'))} руб."))
    elems.append(P(f"Стоимость работ по текущему ремонту, выполненных за отчетный период: {money(report.get('repair_work_cost'))} руб."))
    elems.append(P(f"Остаток (перерасход, сальдо) на 31 декабря отчетного периода: {balance_text(report.get('repair_end_balance'))}"))
    data = [["N", "Наименование", "Основание", "Стоимость", "Объем", "Акт/ссылка"]]
    total_repair = 0.0
    for idx, item in enumerate(bundle["repair"], 1):
        total_repair += to_float(item.get("cost"))
        data.append([idx, item.get("name", ""), item.get("basis", ""), money(item.get("cost")), item.get("volume", ""), item.get("act_details", "")])
    data.append(["ИТОГО", "", "", money(total_repair), "", ""])
    add_table(data, [10*mm, 48*mm, 34*mm, 24*mm, 28*mm, 42*mm])

    elems.append(P("3. Стоимость услуг по управлению многоквартирным домом, оказанных за отчетный период:", section))
    elems.append(P(f"{money(_management_cost(bundle))} руб."))

    elems.append(P("4. Сведения о претензионно-исковой работе:", section))
    claims = bundle["claims"] or {}
    add_table([["N", "Претензий", "Исков / судебных приказов", "Поступило, руб."], [1, claims.get("pretensions_count", 0), claims.get("lawsuits_count", 0), money(claims.get("recovered_amount", 0))]], [12*mm, 50*mm, 74*mm, 50*mm])

    elems.append(P("5. Сведения о начислениях и поступлении средств за отчетный период:", section))
    data = [["N", "Вид платежа", "Задолженность на начало", "Начислено", "Поступило", "Задолженность на 1 января следующего периода"]]
    total = {"debt_start": 0.0, "accrued": 0.0, "paid": 0.0, "debt_end": 0.0}
    for idx, fin in enumerate(_financial_rows(bundle), 1):
        name = "Платежи собственников помещений в многоквартирном доме" if fin.get("payer_type") == "Собственники" else "Платежи нанимателей помещений в многоквартирном доме"
        for key in total:
            total[key] += to_float(fin.get(key))
        data.append([idx, name, money(fin.get("debt_start")), money(fin.get("accrued")), money(fin.get("paid")), money(fin.get("debt_end"))])
    data.append(["ИТОГО", "", money(total["debt_start"]), money(total["accrued"]), money(total["paid"]), money(total["debt_end"])])
    add_table(data, [10*mm, 62*mm, 30*mm, 28*mm, 28*mm, 28*mm])
    doc.build(elems)
    return out_path
